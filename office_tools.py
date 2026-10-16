"""Word, Excel and PowerPoint files, made from structured content (the `make` tool's docx, xlsx and pptx kinds).

Like the PDF engine: the model gives content, the engine owns the layout. The file is saved in the workspace with the exact spec
beside it (so a later turn can read it, change it and make the file again) and also stored for a download link. Content is read
forgivingly (a plain string is a paragraph, a bare list is bullets) because small models rarely get nested structure exactly right.
"""
import io
import json
import os
import re

import backup_tools
import file_tools
import generated_files

_MAX_BLOCKS = 500
_MAX_SHEETS, _MAX_ROWS, _MAX_COLS = 20, 5000, 50
_MAX_SLIDES = 60
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")
_EXT = {"docx": ".docx", "xlsx": ".xlsx", "pptx": ".pptx"}


def _text(value):
    return "" if value is None else str(value)


def _stem(args, default):
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", _text(args.get("filename")).rsplit(".", 1)[0]).strip("_")[:60]
    return stem or default


# ---------- Word ----------

def _build_docx(args):
    from docx import Document
    blocks = args.get("blocks") or args.get("content") or args.get("sections")
    if isinstance(blocks, str):
        blocks = [blocks]
    if not isinstance(blocks, list) or not blocks:
        raise ValueError("blocks must be a list: each a heading, paragraph, bullets, numbered, table, quote or page_break (or just text).")
    doc = Document()
    if args.get("title"):
        doc.add_heading(_text(args["title"]), 0)
    for block in blocks[:_MAX_BLOCKS]:
        if isinstance(block, str):
            doc.add_paragraph(block)
            continue
        if isinstance(block, list):
            block = {"type": "bullets", "items": block}
        kind = _text(block.get("type") or "paragraph").lower()
        if kind == "heading":
            doc.add_heading(_text(block.get("text")), min(max(int(block.get("level") or 1), 1), 4))
        elif kind in ("bullets", "bullet", "list"):
            for item in block.get("items") or []:
                doc.add_paragraph(_text(item), style="List Bullet")
        elif kind in ("numbered", "ordered"):
            for item in block.get("items") or []:
                doc.add_paragraph(_text(item), style="List Number")
        elif kind == "table":
            header = [_text(h) for h in block.get("header") or []]
            rows = block.get("rows") or []
            width = max([len(header)] + [len(r) for r in rows if isinstance(r, list)] or [1])
            table = doc.add_table(rows=1 if header else 0, cols=width)
            table.style = "Light Grid Accent 1"
            if header:
                for i, h in enumerate(header):
                    table.rows[0].cells[i].text = h
            for row in rows[:_MAX_ROWS]:
                cells = table.add_row().cells
                for i, value in enumerate(row if isinstance(row, list) else [row]):
                    if i < width:
                        cells[i].text = _text(value)
        elif kind == "quote":
            doc.add_paragraph(_text(block.get("text")), style="Quote")
        elif kind in ("page_break", "pagebreak"):
            doc.add_page_break()
        else:
            doc.add_paragraph(_text(block.get("text")))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), {"blocks": min(len(blocks), _MAX_BLOCKS)}


# ---------- Excel ----------

def _build_xlsx(args):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    sheets = args.get("sheets")
    if sheets is None and (args.get("rows") or args.get("header")):
        sheets = [{"name": args.get("name") or "Sheet1", "header": args.get("header"), "rows": args.get("rows")}]
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("sheets must be a list of {name, header, rows}: header is a list of column names, rows a list of lists.")
    wb = Workbook()
    wb.remove(wb.active)
    for number, sheet in enumerate(sheets[:_MAX_SHEETS], 1):
        ws = wb.create_sheet((_SHEET_BAD.sub(" ", _text(sheet.get("name")) or f"Sheet{number}").strip() or f"Sheet{number}")[:31])
        header = sheet.get("header") or []
        rows = sheet.get("rows") or []
        if header:
            ws.append([_text(h) for h in header[:_MAX_COLS]])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4F46E5")
            ws.freeze_panes = "A2"
        for row in rows[:_MAX_ROWS]:
            ws.append([v if isinstance(v, (int, float, bool)) or v is None else _text(v) for v in (row if isinstance(row, list) else [row])[:_MAX_COLS]])
        widths = sheet.get("widths")
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            if isinstance(widths, list) and col <= len(widths) and isinstance(widths[col - 1], (int, float)):
                ws.column_dimensions[letter].width = widths[col - 1]
            else:
                longest = max((len(_text(c.value)) for c in ws[letter][:200]), default=8)
                ws.column_dimensions[letter].width = min(max(longest + 2, 8), 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), {"sheets": len(wb.sheetnames), "rows": sum(max(ws.max_row - (1 if ws.freeze_panes else 0), 0) for ws in wb.worksheets)}


# ---------- PowerPoint ----------

def _widen(slide):
    """The default template is 4:3; slides here are 16:9, so the placeholders are stretched to match."""
    for shape in slide.shapes:
        if shape.left is not None and shape.width is not None:
            shape.left, shape.width = int(shape.left * 4 / 3), int(shape.width * 4 / 3)


def _build_pptx(args):
    from pptx import Presentation
    from pptx.util import Inches
    slides = args.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("slides must be a list of {title, bullets, notes}: bullets a list of short lines (a nested list makes sub-points).")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    if args.get("title"):
        first = prs.slides.add_slide(prs.slide_layouts[0])
        first.shapes.title.text = _text(args["title"])
        if len(first.placeholders) > 1:
            first.placeholders[1].text = _text(args.get("subtitle"))
        _widen(first)
    for item in slides[:_MAX_SLIDES]:
        item = item if isinstance(item, dict) else {"title": _text(item)}
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = _text(item.get("title"))
        body = slide.placeholders[1].text_frame
        lines = item.get("bullets") or ([item["text"]] if item.get("text") else [])
        first_line = True
        for line in lines[:12]:
            for level, content in ((0, line),) if not isinstance(line, list) else [(1, sub) for sub in line]:
                para = body.paragraphs[0] if first_line else body.add_paragraph()
                first_line = False
                para.text = _text(content)
                para.level = level
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = _text(item["notes"])
        _widen(slide)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), {"slides": len(prs.slides)}


_BUILDERS = {"docx": _build_docx, "xlsx": _build_xlsx, "pptx": _build_pptx}

SPEC_SCHEMAS = {
    "docx": {"description": "A Word document: a title and a list of blocks, in order.",
             "parameters": {"type": "object", "required": ["blocks"], "properties": {
                 "filename": {"type": "string", "description": "Name without extension."}, "title": {"type": "string"},
                 "blocks": {"type": "array", "description": "Each: {type: heading, text, level 1-4} | {type: paragraph, text} | {type: bullets, items: [..]} | {type: numbered, items: [..]} | "
                                                            "{type: table, header: [..], rows: [[..]]} | {type: quote, text} | {type: page_break}. A plain string is a paragraph.",
                            "items": {}}}}},
    "xlsx": {"description": "An Excel workbook: one or more sheets, each a header row and rows of cells. A cell that starts with = is a formula.",
             "parameters": {"type": "object", "required": ["sheets"], "properties": {
                 "filename": {"type": "string", "description": "Name without extension."},
                 "sheets": {"type": "array", "description": "Each: {name, header: [column names], rows: [[cells]], widths: [optional column widths]}.", "items": {}}}}},
    "pptx": {"description": "A PowerPoint deck (16:9): an optional title slide, then one slide per item.",
             "parameters": {"type": "object", "required": ["slides"], "properties": {
                 "filename": {"type": "string", "description": "Name without extension."}, "title": {"type": "string", "description": "Makes a title slide first."}, "subtitle": {"type": "string"},
                 "slides": {"type": "array", "description": "Each: {title, bullets: [short lines; a nested list makes sub-points], notes: speaker notes}.", "items": {}}}}},
}


def generate(kind, args, workspace=""):
    args = args if isinstance(args, dict) else {}
    try:
        data, summary = _BUILDERS[kind](args)
    except ValueError as e:
        return {"error": str(e)}
    except Exception as e:
        return {"error": f"Could not build the {kind} file: {e}"}
    stem = _stem(args, {"docx": "document", "xlsx": "workbook", "pptx": "presentation"}[kind])
    filename = stem + _EXT[kind]
    result = {"filename": filename, "url": generated_files.url_for(generated_files.store(data, filename)), "size": len(data), **summary}
    if workspace:
        spec = {**{k: v for k, v in args.items() if k != "filename"}, "filename": stem}
        try:
            for rel, content in ((filename, data), (filename + ".json", (json.dumps(spec, indent=2, ensure_ascii=False) + "\n").encode("utf-8"))):
                path = file_tools.resolve_workspace_path(workspace, rel)
                if os.path.exists(path):
                    backup_tools.backup_file(path)
                with open(path, "wb") as f:
                    f.write(content)
            result["saved_to"] = filename
            result["spec"] = filename + ".json"
            result["revise"] = f"To change it: read_file {filename}.json, edit the content, and make again with filename \"{stem}\". That replaces both files."
        except (OSError, ValueError) as e:
            result["workspace_save_error"] = f"The download link works, but the files could not be saved in the workspace: {e}"
    return result
