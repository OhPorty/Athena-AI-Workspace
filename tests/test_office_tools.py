"""make kind docx / xlsx / pptx: built from structured content, saved with their spec, readable back."""
import io
import json

import pytest

import office_tools
import surface
from test_chat_loop_recovery import call, run_turn, text_round, tool_round


def read(path):
    return open(path, "rb").read()


def test_a_word_document_has_exactly_the_content_it_was_given(tmp_path):
    out = office_tools.generate("docx", {"filename": "Quarterly Report!", "title": "Q3", "blocks": [
        {"type": "heading", "text": "Summary", "level": 1}, "A plain string is a paragraph.", ["bullet one", "bullet two"],
        {"type": "numbered", "items": ["first", "second"]}, {"type": "table", "header": ["Item", "Cost"], "rows": [["Rent", 1200], ["Food", 400]]},
        {"type": "quote", "text": "Measure twice."}, {"type": "page_break"}, {"type": "paragraph", "text": "The end."}]}, str(tmp_path))
    assert out["filename"] == "Quarterly_Report.docx" and out["blocks"] == 8 and out["url"].startswith("/api/generated-files/")
    from docx import Document
    doc = Document(str(tmp_path / "Quarterly_Report.docx"))
    texts = [p.text for p in doc.paragraphs]
    assert texts[:3] == ["Q3", "Summary", "A plain string is a paragraph."] and "bullet two" in texts and "Measure twice." in texts and texts[-1] == "The end."
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Item", "Cost"] and doc.tables[0].rows[2].cells[1].text == "400"
    assert [p.style.name for p in doc.paragraphs if p.text in ("bullet one", "first")] == ["List Bullet", "List Number"]


def test_an_excel_workbook_keeps_types_formulas_and_sheet_names(tmp_path):
    out = office_tools.generate("xlsx", {"filename": "budget", "sheets": [
        {"name": "Budget: 2026/Q3", "header": ["Item", "Cost", "Paid"], "rows": [["Rent", 1200, True], ["Food", 400, False], ["Total", "=SUM(B2:B3)", None]]},
        {"header": ["Only"], "rows": [[1], [2]]}]}, str(tmp_path))
    assert out["sheets"] == 2 and out["rows"] == 5
    from openpyxl import load_workbook
    wb = load_workbook(str(tmp_path / "budget.xlsx"))
    assert wb.sheetnames == ["Budget  2026 Q3", "Sheet2"]                       # characters Excel forbids are replaced, a missing name is filled in
    ws = wb["Budget  2026 Q3"]
    assert ws["A1"].font.bold and ws.freeze_panes == "A2" and ws["B2"].value == 1200 and ws["C2"].value is True and ws["B4"].value == "=SUM(B2:B3)"
    assert ws.column_dimensions["A"].width >= 8


def test_a_powerpoint_deck_is_16_by_9_with_bullets_subpoints_and_notes(tmp_path):
    out = office_tools.generate("pptx", {"filename": "deck", "title": "My deck", "subtitle": "2026", "slides": [
        {"title": "Why", "bullets": ["first", ["detail a", "detail b"], "second"], "notes": "say this"}, {"title": "Close", "text": "Thanks"}]}, str(tmp_path))
    assert out["slides"] == 3
    from pptx import Presentation
    prs = Presentation(str(tmp_path / "deck.pptx"))
    assert round(prs.slide_width / prs.slide_height, 2) == 1.78
    assert prs.slides[0].shapes.title.text == "My deck" and prs.slides[0].placeholders[1].text == "2026"
    body = prs.slides[1].placeholders[1].text_frame
    assert [(p.text, p.level) for p in body.paragraphs] == [("first", 0), ("detail a", 1), ("detail b", 1), ("second", 0)]
    assert prs.slides[1].notes_slide.notes_text_frame.text == "say this" and prs.slides[2].placeholders[1].text_frame.text == "Thanks"


def test_the_spec_is_saved_beside_the_file_so_a_later_turn_can_revise_it(tmp_path):
    out = office_tools.generate("docx", {"filename": "memo", "blocks": ["one"]}, str(tmp_path))
    assert out["saved_to"] == "memo.docx" and out["spec"] == "memo.docx.json" and "memo.docx.json" in out["revise"]
    assert json.loads((tmp_path / "memo.docx.json").read_text()) == {"blocks": ["one"], "filename": "memo"}
    office_tools.generate("docx", {"filename": "memo", "blocks": ["one", "two"]}, str(tmp_path))                 # making it again replaces both
    from docx import Document
    assert [p.text for p in Document(str(tmp_path / "memo.docx")).paragraphs] == ["one", "two"]
    assert "saved_to" not in office_tools.generate("docx", {"blocks": ["x"]}, "")                                # no workspace: a download link only


def test_missing_or_empty_content_says_what_is_needed(tmp_path):
    assert "blocks" in office_tools.generate("docx", {}, str(tmp_path))["error"]
    assert "sheets" in office_tools.generate("xlsx", {"sheets": []}, str(tmp_path))["error"]
    assert "slides" in office_tools.generate("pptx", {"title": "x"}, str(tmp_path))["error"]
    assert office_tools.generate("xlsx", {"header": ["a"], "rows": [[1]]}, str(tmp_path))["sheets"] == 1       # loose top-level fields count as one sheet


def test_loose_kind_names_reach_the_right_engine():
    for word, kind in (("word", "docx"), ("Word document", "docx"), ("excel", "xlsx"), ("spreadsheet", "xlsx"), ("powerpoint", "pptx"), ("slides", "pptx"), ("presentation", "pptx"),
                       ("resume", "pdf"), ("pdf", "pdf"), ("background music", "music")):
        assert surface._kind_of(word) == kind, word
    assert surface.to_engine("make", {"kind": "powerpoint", "spec": {"slides": [{"title": "a"}]}})[0] == "generate_pptx"


@pytest.mark.parametrize("kind", ["docx", "xlsx", "pptx"])
def test_the_guides_own_example_actually_builds(kind, tmp_path):
    guide = surface.make_guide(kind)
    assert guide["kind"] == kind and guide["spec_fields"]["required"] and "not the shell" in guide["how_to_call"]
    out = office_tools.generate(kind, guide["example_call"]["spec"], str(tmp_path))
    assert "error" not in out and out["size"] > 1000
    assert surface.to_engine("make", {"kind": kind})[0] == "_surface_error"                                    # no spec: shown the guide first


def test_make_a_deck_in_a_real_turn_and_it_is_snapshotted_first(monkeypatch, tmp_path):
    events, seen, saved, ws = run_turn(monkeypatch, tmp_path, [tool_round(call("make", kind="pptx", spec={"slides": [{"title": "Hello", "bullets": ["one"]}]})), text_round("done")],
                                       files={"notes.txt": "keep"})
    result = json.loads([m for m in seen[1] if m["role"] == "tool"][0]["content"])
    assert result["saved_to"] == "presentation.pptx" and (ws / "presentation.pptx").exists()
    assert [e for e in events if e.get("type") == "snapshot"]                                                  # a change was about to happen: the copy was saved first
