"""Non-text files: reading them back (so a model can check its own work), downloading them
(so the user can get them), and the sandbox settings that let a script make them."""
import zipfile

import pytest
from fastapi.responses import FileResponse, JSONResponse

import bash_tools
import binary_files
import file_tools


@pytest.fixture
def ws(tmp_path):
    return tmp_path


def _docx(path):
    import docx
    d = docx.Document()
    d.add_heading("Quarterly report", 1)
    d.add_paragraph("Revenue grew.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Q", "Rev"
    d.save(path)


def test_docx_xlsx_pptx_pdf_zip_and_images_read_back_as_text(ws):
    import openpyxl, pypdf
    from PIL import Image
    from pptx import Presentation
    _docx(ws / "r.docx")
    wb = openpyxl.Workbook(); wb.active.title = "Sales"; wb.active.append(["item", "qty"]); wb.active.append(["apple", 3]); wb.save(ws / "s.xlsx")
    p = Presentation(); p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = "Launch plan"; p.save(ws / "d.pptx")
    w = pypdf.PdfWriter(); w.add_blank_page(100, 100); w.write(ws / "b.pdf")
    Image.new("RGB", (64, 32)).save(ws / "i.png")
    with zipfile.ZipFile(ws / "z.zip", "w") as z:
        z.writestr("readme.txt", "hi")
    got = {n: file_tools.read_file(str(ws), n) for n in ("r.docx", "s.xlsx", "d.pptx", "b.pdf", "i.png", "z.zip")}
    assert "Quarterly report" in got["r.docx"]["content"] and "Q | Rev" in got["r.docx"]["content"]
    assert "[sheet Sales]" in got["s.xlsx"]["content"] and "apple\t3" in got["s.xlsx"]["content"]
    assert "Launch plan" in got["d.pptx"]["content"] and "[page 1]" in got["b.pdf"]["content"]
    assert "64x32" in got["i.png"]["content"] and "readme.txt" in got["z.zip"]["content"]
    assert all(r["binary"] and "error" not in r for r in got.values()) and got["i.png"]["kind"] == "png"


def test_plain_text_is_unchanged_and_other_binaries_say_so_instead_of_returning_garbage(ws):
    (ws / "a.py").write_text("print(1)\n")
    (ws / "x.bin").write_bytes(b"\x00\x01\x02" * 100)
    text = file_tools.read_file(str(ws), "a.py")
    assert text["content"] == "print(1)\n" and "binary" not in text
    blob = file_tools.read_file(str(ws), "x.bin")
    assert blob["binary"] and "no text view" in blob["content"] and "\x00" not in blob["content"]


def test_a_damaged_office_file_is_described_not_raised(ws):
    (ws / "bad.docx").write_bytes(b"this is not a zip")
    out = file_tools.read_file(str(ws), "bad.docx")
    assert "error" not in out and "damaged" in out["content"] and out["binary"]


def test_download_serves_workspace_files_and_only_images_display_inline(ws):
    (ws / "r.docx").write_bytes(b"PK")
    (ws / "p.png").write_bytes(b"\x89PNG")
    (ws / "page.html").write_text("<script>alert(1)</script>")
    (ws / "logo.svg").write_text("<svg onload='alert(1)'/>")
    disposition = lambda name, inline: file_tools.download_workspace_file(str(ws), name, inline).headers["content-disposition"]
    assert isinstance(file_tools.download_workspace_file(str(ws), "r.docx"), FileResponse)
    assert disposition("p.png", True).startswith("inline") and disposition("p.png", False).startswith("attachment")
    assert disposition("page.html", True).startswith("attachment") and disposition("logo.svg", True).startswith("attachment")


def test_download_refuses_escapes_and_missing_files(ws):
    for path in ("../../etc/passwd", "/etc/passwd"):
        r = file_tools.download_workspace_file(str(ws), path)
        assert isinstance(r, JSONResponse) and r.status_code == 400
    assert file_tools.download_workspace_file(str(ws), "nope.docx").status_code == 404
    assert file_tools.download_workspace_file("", "a").status_code == 400


def test_sandbox_installs_into_the_workspace_so_they_survive_between_calls(ws):
    cmd = bash_tools._build_docker_sandbox_cmd("pip", ["install", "x"], str(ws))
    assert f"PYTHONUSERBASE={ws}/.athena/pylibs" in cmd
    assert "--read-only" in cmd   # the root filesystem stays locked; only the workspace is writable
    assert {"zip", "unzip"} <= bash_tools.BASH_EXEC_ALLOWED_COMMANDS


def test_tool_description_names_the_libraries_that_are_actually_installed():
    desc = next(t["function"]["description"] for t in bash_tools.BASH_TOOL_SCHEMAS if t["function"]["name"] == "bash_exec")
    dockerfile = open(bash_tools.os.path.join(bash_tools.os.path.dirname(bash_tools.__file__), "sandbox", "Dockerfile")).read().lower()
    for lib in ("python-docx", "openpyxl", "python-pptx", "reportlab", "pillow", "matplotlib", "cairosvg"):
        assert lib in dockerfile and lib.lower().replace("pillow", "pillow") in desc.lower(), lib


def test_audio_displays_inline_and_supports_seeking_ranges(ws):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    import wave
    with wave.open(str(ws / "tune.wav"), "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(8000); w.writeframes(b"\x00\x01" * 8000)
    (ws / "page.html").write_text("<script>alert(1)</script>")
    app = FastAPI(); app.include_router(file_tools.router)
    client = TestClient(app)
    url = "/api/workspace/download"
    r = client.get(url, params={"workspace": str(ws), "path": "tune.wav", "inline": "true"})
    assert r.status_code == 200 and r.headers["content-disposition"].startswith("inline") and r.headers["content-type"].startswith("audio/")
    assert r.headers["accept-ranges"] == "bytes"   # browsers need this to seek or scrub a long track
    part = client.get(url, params={"workspace": str(ws), "path": "tune.wav", "inline": "true"}, headers={"Range": "bytes=0-99"})
    assert part.status_code == 206 and len(part.content) == 100
    assert client.get(url, params={"workspace": str(ws), "path": "page.html", "inline": "true"}).headers["content-disposition"].startswith("attachment")


# ---------- audio: a model can't listen, so read_file measures ----------

def _write_wav(path, samples, rate=44100):
    import numpy as np, wave
    with wave.open(str(path), "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(rate)
        w.writeframes((np.clip(samples, -1, 1) * 32767).astype("<i2").tobytes())


def test_the_audio_report_calls_noise_noise_and_music_music(ws):
    import numpy as np
    rate, t = 44100, np.arange(44100 * 12) / 44100
    rng = np.random.default_rng(3)
    hiss = rng.normal(0, 0.1, len(t)) + (rng.random(len(t)) < 0.0005) * 0.8       # the failed lofi: crackle and hiss only
    chords = sum(np.sin(2 * np.pi * f * t) for f in (220, 261.6, 329.6)) * 0.2
    beat = np.zeros_like(t); beat[(t % 0.5) < 0.05] = 1; chords = chords * (0.4 + 0.6 * beat)   # pulsing chord
    _write_wav(ws / "noise.wav", hiss)
    _write_wav(ws / "music.wav", chords)
    noise = file_tools.read_file(str(ws), "noise.wav")["content"]
    music = file_tools.read_file(str(ws), "music.wav")["content"]
    assert "noise or hiss, not music" in noise and "none (noise-like)" in noise and "layer" in noise
    assert "Pitched notes: yes" in music and "Problems: none found" in music
    assert file_tools.read_file(str(ws), "music.wav")["kind"] == "wav"


def test_the_audio_report_flags_silence_clipping_and_whisper_quiet_files(ws):
    import numpy as np
    t = np.arange(44100 * 3) / 44100
    tone = np.sin(2 * np.pi * 440 * t)
    _write_wav(ws / "silent.wav", tone * 0.0001)
    _write_wav(ws / "clipped.wav", np.clip(tone * 3, -1, 1))
    _write_wav(ws / "quiet.wav", tone * 0.012)
    reports = {n: file_tools.read_file(str(ws), n + ".wav")["content"] for n in ("silent", "clipped", "quiet")}
    assert "nearly silent" in reports["silent"] and "clipping" in reports["clipped"] and "very quiet" in reports["quiet"]
    assert "no regular beat" in reports["quiet"]   # a steady tone has no beat


def test_a_stereo_wav_and_other_audio_formats_do_not_break_read_file(ws):
    import numpy as np, wave
    with wave.open(str(ws / "st.wav"), "wb") as w:
        w.setnchannels(2); w.setsampwidth(2); w.setframerate(44100)
        x = (np.sin(2 * np.pi * 330 * np.arange(44100 * 2) / 44100) * 8000).astype("<i2")
        w.writeframes(np.stack([x, x], 1).tobytes())
    (ws / "song.mp3").write_bytes(b"ID3" + b"\x00" * 200)
    assert "stereo" in file_tools.read_file(str(ws), "st.wav")["content"]
    assert file_tools.read_file(str(ws), "song.mp3")["binary"]


def test_the_audio_starter_runs_and_makes_music_by_the_audio_reports_own_standard(tmp_path):
    import scaffold_tools, subprocess, sys
    r = scaffold_tools.scaffold(str(tmp_path), {"kind": "audio"})
    assert sorted(r["created"]) == ["song.py", "synth.py"] and r["open"] == "song.py"
    done = subprocess.run([sys.executable, "song.py"], cwd=tmp_path, capture_output=True, text=True, timeout=120)
    assert done.returncode == 0, done.stderr
    report = file_tools.read_file(str(tmp_path), "song.wav")["content"]
    assert "Pitched notes: yes" in report and "Problems: none found" in report


def test_the_synth_toolkit_refuses_a_sound_in_the_wrong_units(tmp_path):
    import importlib.util, numpy as np, scaffold_tools
    scaffold_tools.scaffold(str(tmp_path), {"kind": "audio"})
    spec = importlib.util.spec_from_file_location("synth_under_test", tmp_path / "synth.py")
    synth = importlib.util.module_from_spec(spec); spec.loader.exec_module(synth)
    track = synth.Track(2)
    with pytest.raises(ValueError, match="-1..1"):
        track.add(np.random.default_rng(0).normal(0, 100, 4000), 0)   # raw 16-bit-scale noise, the exact mistake that buried a whole song
    assert synth.note("C4") == 60 and synth.note("Bb2") == 46 and synth.chord("D3", "m7") == [50, 53, 57, 60]
    assert all(np.abs(f()).max() <= 1.0 + 1e-9 for f in (synth.kick, synth.snare, synth.hat))


# ---------- grep patterns are not shell syntax ----------

@pytest.mark.parametrize("arg", [r"master\|_mix_into", "foo|bar", "=>", "->", "a<b", "def x|return y"])
def test_regex_and_arrow_text_in_arguments_is_allowed(arg):
    assert bash_tools._validate_readonly_stage("grep", ["-n", arg, "a.py"]) is None


@pytest.mark.parametrize("arg", ["|", "a | b", "&&", "x && y", ">out.txt", "2>/dev/null", "2>&1", "$(id)", "`id`", ";"])
def test_real_shell_syntax_is_still_refused(arg):
    assert "shell operator" in bash_tools._validate_readonly_stage("grep", ["-n", arg, "a.py"])["error"]
    assert "shell operator" in bash_tools._validate_bash_exec_call("python3", ["x.py", arg])["error"]
